"""Read teammate workbook locally without inferring event categories or intervals.

Detailed descriptions are retained only in the private output records. The CLI
prints aggregate counts and issue codes, never the descriptions themselves.
"""

import hashlib
import math
import re
from pathlib import Path


def point_seconds(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        result = float(value)
    else:
        text = str(value).lower().strip()
        if re.search(r"\d\s*[-–—]\s*\d", text):
            raise ValueError("interval_requires_explicit_event_import")
        if re.fullmatch(r"\d+(?::\d+(?:\.\d+)?){1,2}", text):
            result = 0.0
            for part in text.split(":"):
                result = result * 60 + float(part)
        else:
            minutes = re.search(r"(\d+)\s*м", text)
            seconds = re.search(r"(\d+)\s*с", text)
            if not minutes and not seconds:
                raise ValueError("unrecognized_point_timecode")
            result = float(
                (int(minutes[1]) * 60 if minutes else 0)
                + (int(seconds[1]) if seconds else 0)
            )
    if not math.isfinite(result) or result < 0:
        raise ValueError("invalid_point_timecode")
    return result


def explicit_binary(text):
    text = str(text or "").strip().lower()
    if re.fullmatch(r"без\s+нарушени[а-яё]*[.!\s]*", text):
        return 0
    if re.fullmatch(r"нарушени[а-яё]*[.!\s]*", text):
        return 1
    return None


def import_rows(rows, baseline, source_name, source_sha256, sheet):
    known = {}
    for entry in baseline:
        stem = Path(entry["video_id"]).stem
        if stem in known:
            raise ValueError("ambiguous_baseline_filename")
        if entry["binary_label"] not in (0, 1):
            raise ValueError("invalid_baseline_binary_label")
        known[stem] = entry
    records = []
    seen = set()
    for row_number, row in enumerate(rows, 2):
        name, raw_timecode, annotation = (list(row[:3]) + [None] * 3)[:3]
        if all(value is None for value in (name, raw_timecode, annotation)):
            continue
        if name is None:
            raise ValueError(f"row_{row_number}:missing_video_id")
        name = str(name).strip()
        if Path(name).name != name or name in (".", ".."):
            raise ValueError(f"row_{row_number}:invalid_video_id")
        # Match exact Unicode stems. Latin X and Cyrillic Х are NEVER merged.
        reference = known.get(Path(name).stem)
        if reference is None:
            raise ValueError(f"row_{row_number}:no_exact_filename_match")
        video_id = reference["video_id"]
        if video_id in seen:
            raise ValueError(f"row_{row_number}:duplicate_video_id")
        seen.add(video_id)
        label = reference["binary_label"]
        annotation = "" if annotation is None else str(annotation)
        stated = explicit_binary(annotation)
        if stated is not None and stated != label:
            raise ValueError(f"row_{row_number}:binary_label_conflict")
        timestamp = point_seconds(raw_timecode)
        duration = reference["duration_sec"]
        if not math.isfinite(duration) or duration <= 0:
            raise ValueError("invalid_baseline_duration")
        if timestamp is not None and timestamp > duration:
            raise ValueError(f"row_{row_number}:timecode_outside_video")
        issues = []
        if timestamp != reference.get("timecode_sec"):
            issues.append("reference_timecode_changed")
        if label and timestamp is None:
            issues.append("missing_reference_timecode")
        if label and (not annotation.strip() or stated is not None):
            issues.append("missing_detailed_annotation")
        if label:
            issues.extend(
                [
                    "category_not_structured",
                    "event_interval_missing",
                    "rule_not_verified",
                ]
            )
        records.append(
            {
                "video_id": video_id,
                "binary_label": label,
                "reference_timecode_sec": timestamp,
                "raw_timecode": raw_timecode,
                "team_annotation": annotation,
                "category": None,
                "start_sec": None,
                "end_sec": None,
                "observed_action": None,
                "rule_id": None,
                "scene_id": None,
                "split": None,
                "origin": "open_dataset",
                "normative_source": None,
                "source_verification_status": "unverified",
                "ready_for_event_training": False,
                "issues": issues,
                "provenance": {
                    "file": source_name,
                    "sha256": source_sha256,
                    "sheet": sheet,
                    "row": row_number,
                    "video_cell": f"A{row_number}",
                    "timecode_cell": f"B{row_number}",
                    "annotation_cell": f"C{row_number}",
                },
            }
        )
    return records


def summarize(records):
    positive = [row for row in records if row["binary_label"] == 1]
    return {
        "video_count": len(records),
        "positive_video_count": len(positive),
        "negative_video_count": sum(row["binary_label"] == 0 for row in records),
        "positive_with_description": sum(
            bool(row["team_annotation"].strip())
            and "missing_detailed_annotation" not in row["issues"]
            for row in positive
        ),
        "positive_with_reference_timecode": sum(
            row["reference_timecode_sec"] is not None for row in positive
        ),
        "reference_timecode_changes": sum(
            "reference_timecode_changed" in row["issues"] for row in records
        ),
        "structured_category_count": sum(
            row["category"] is not None for row in positive
        ),
        "complete_event_interval_count": sum(
            row["start_sec"] is not None and row["end_sec"] is not None
            for row in positive
        ),
        "ready_for_event_training_count": sum(
            row["ready_for_event_training"] for row in records
        ),
        "scene_groups_assigned": sum(row["scene_id"] is not None for row in records),
        "note": "Point annotations and teammate descriptions preserved; no categories, intervals, rules or splits inferred. No training started.",
    }


def read_workbook(path, baseline, sheet="labels"):
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise ValueError("sheet_not_found")
        rows = list(book[sheet].values)
        headers = (
            [str(value or "").strip().lower() for value in rows[0][:3]] if rows else []
        )
        if headers != ["наименование файла", "таймстемп", "нарушение / без нарушения"]:
            raise ValueError("unexpected_workbook_headers")
        return import_rows(
            rows[1:],
            baseline,
            Path(path).name,
            hashlib.sha256(Path(path).read_bytes()).hexdigest(),
            sheet,
        )
    finally:
        book.close()
